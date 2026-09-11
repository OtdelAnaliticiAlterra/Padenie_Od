from datetime import date
from dateutil.relativedelta import relativedelta

import pandas as pd
from openpyxl.utils import get_column_letter

import config
from db import get_engine

from db import run_query  # вместо get_engine


def _fmt(d: date) -> str:
    """Формат даты для SQL-параметра."""
    return d.strftime("%Y-%m-%d 00:00:00.000")


def _label(d1: date, d2: date) -> str:
    """Человеческая подпись периода; d2 — эксклюзивная граница."""
    return f"{d1:%d.%m.%Y} - {(d2 - relativedelta(days=1)):%d.%m.%Y}"


def build_period_params() -> dict:
    """Считает все 6 дат и 3 подписи периодов на основе config.CURRENT_START."""
    y, m, d = config.CURRENT_START
    cur_start = date(y, m, d)
    cur_end = cur_start + relativedelta(months=config.PERIOD_MONTHS)

    prev_start = cur_start - relativedelta(years=1)
    prev_end = cur_end - relativedelta(years=1)

    fut_start = prev_end
    fut_end = fut_start + relativedelta(months=config.PERIOD_MONTHS)

    return {
        "cur_start":  _fmt(cur_start),
        "cur_end":    _fmt(cur_end),
        "prev_start": _fmt(prev_start),
        "prev_end":   _fmt(prev_end),
        "fut_start":  _fmt(fut_start),
        "fut_end":    _fmt(fut_end),
        "period_label": _label(cur_start, cur_end),
        "prev_label":   _label(prev_start, prev_end),
        "future_label": _label(fut_start, fut_end),
    }


def fetch_data(params: dict) -> pd.DataFrame:
    sql = config.SQL_FILE.read_text(encoding="utf-8")
    result = run_query(sql, params)
    # Приводим к DataFrame
    df = pd.DataFrame(result.fetchall(), columns=result.keys())
    return df


def save_excel(df: pd.DataFrame, params: dict) -> str:
    filename = config.OUTPUT_FILENAME_TEMPLATE.format(
        start=params["period_label"].split(" - ")[0].replace(".", ""),
        end=params["period_label"].split(" - ")[1].replace(".", ""),
    )
    out_path = config.OUTPUT_DIR / filename

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Отчет", index=False)
        ws = writer.sheets["Отчет"]

        # Автоширина
        for i, col in enumerate(df.columns, start=1):
            max_len = max(
                len(str(col)),
                df[col].astype(str).str.len().max() if len(df) else 0,
            )
            ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 40)

        # Форматы чисел
        for col in config.MONEY_COLUMNS:
            if col in df.columns:
                idx = df.columns.get_loc(col) + 1
                for row in range(2, len(df) + 2):
                    ws.cell(row=row, column=idx).number_format = '#,##0.00'

        for col in config.PERCENT_COLUMNS:
            if col in df.columns:
                idx = df.columns.get_loc(col) + 1
                for row in range(2, len(df) + 2):
                    ws.cell(row=row, column=idx).number_format = '0.00"%"'

        # Шапка + фильтр
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    return str(out_path)


def main():
    params = build_period_params()
    print("Периоды:")
    for k in ("period_label", "prev_label", "future_label"):
        print(f"  {k}: {params[k]}")

    print("Выполняю запрос...")
    df = fetch_data(params)
    print(f"Получено строк: {len(df)}")

    path = save_excel(df, params)
    print(f"Готово: {path}")


if __name__ == "__main__":
    main()